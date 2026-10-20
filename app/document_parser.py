"""Text extraction for PDF, DOCX, and TXT files.

These functions take raw file bytes and know nothing about HTTP or FastAPI,
so they can be tested and reused on their own. Any problem with the file
content is raised as DocumentParseError with a message safe to show users.
"""

import io
from dataclasses import dataclass

import docx
import pymupdf
from docx.table import Table


class DocumentParseError(Exception):
    """The file could not be read. The message is shown to the user."""


@dataclass
class ExtractedSection:
    # 1-based PDF page number; None when the format has no fixed pages.
    page_number: int | None
    text: str


@dataclass
class ExtractedDocument:
    sections: list[ExtractedSection]
    page_count: int | None = None
    title: str | None = None
    author: str | None = None


def extract_pdf(data: bytes) -> ExtractedDocument:
    """Extract text page by page, keeping each page's number."""
    try:
        pdf = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:  # PyMuPDF raises several error types for bad input
        raise DocumentParseError("The PDF file is damaged or not a valid PDF.") from exc

    with pdf:
        if pdf.needs_pass:
            raise DocumentParseError(
                "The PDF is password-protected. Remove the password and upload it again."
            )

        try:
            sections = [
                ExtractedSection(page_number=index + 1, text=page.get_text())
                for index, page in enumerate(pdf)
            ]
        except Exception as exc:
            raise DocumentParseError(
                "Text could not be read from the PDF's pages."
            ) from exc

        metadata = pdf.metadata or {}
        return ExtractedDocument(
            sections=sections,
            page_count=pdf.page_count,
            title=metadata.get("title") or None,
            author=metadata.get("author") or None,
        )


def extract_docx(data: bytes) -> ExtractedDocument:
    """Extract paragraphs and tables in the order they appear in the document.

    Word works out page breaks only when it lays the document out for display,
    so a DOCX file has no reliable page numbers. The text is one section.
    """
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # BadZipFile, KeyError, ValueError, ... for bad input
        raise DocumentParseError(
            "The DOCX file is damaged or not a valid Word document."
        ) from exc

    blocks: list[str] = []
    for item in document.iter_inner_content():
        if isinstance(item, Table):
            blocks.extend(_table_rows(item))
        elif item.text.strip():
            blocks.append(item.text)

    properties = document.core_properties
    return ExtractedDocument(
        sections=[ExtractedSection(page_number=None, text="\n".join(blocks))],
        title=properties.title or None,
        author=properties.author or None,
    )


def _table_rows(table: Table) -> list[str]:
    """Turn each table row into one line, with cells separated by " | "."""
    rows = []
    for row in table.rows:
        cells: list[str] = []
        for cell in row.cells:
            text = cell.text.strip()
            # python-docx repeats a merged cell's text once per grid column.
            if cells and cells[-1] == text:
                continue
            cells.append(text)
        if any(cells):
            rows.append(" | ".join(cells))
    return rows


def extract_txt(data: bytes) -> ExtractedDocument:
    """Decode a UTF-8 text file (a leading byte-order mark is allowed)."""
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DocumentParseError(
            "The text file is not UTF-8 encoded. Save it as UTF-8 and upload it again."
        ) from exc

    if "\x00" in text:
        raise DocumentParseError("The file looks like binary data, not a text file.")

    text = text.replace("\r\n", "\n")
    return ExtractedDocument(sections=[ExtractedSection(page_number=None, text=text)])
