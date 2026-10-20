import io

import docx
import pytest

from app.document_parser import (
    DocumentParseError,
    extract_docx,
    extract_pdf,
    extract_txt,
)
from tests.conftest import make_docx, make_pdf

# --- PDF ---


def test_pdf_keeps_page_numbers_and_text(two_page_pdf: bytes) -> None:
    result = extract_pdf(two_page_pdf)

    assert result.page_count == 2
    assert [s.page_number for s in result.sections] == [1, 2]
    assert "Alpha page one" in result.sections[0].text
    assert "Beta page two" in result.sections[1].text


def test_pdf_reads_title_and_author(two_page_pdf: bytes) -> None:
    result = extract_pdf(two_page_pdf)

    assert result.title == "Sample"
    assert result.author == "Tester"


def test_pdf_missing_metadata_becomes_none() -> None:
    result = extract_pdf(make_pdf(["text"]))

    assert result.title is None
    assert result.author is None


def test_pdf_blank_page_gives_empty_section() -> None:
    result = extract_pdf(make_pdf(["has text", ""]))

    assert result.sections[1].page_number == 2
    assert result.sections[1].text.strip() == ""


def test_pdf_damaged_raises_parse_error() -> None:
    with pytest.raises(DocumentParseError, match="damaged"):
        extract_pdf(b"%PDF-1.7 this is not really a pdf")


def test_pdf_password_protected_raises_parse_error() -> None:
    with pytest.raises(DocumentParseError, match="password"):
        extract_pdf(make_pdf(["secret"], password="pw"))


# --- DOCX ---


def test_docx_extracts_paragraphs_and_table_in_order(sample_docx: bytes) -> None:
    result = extract_docx(sample_docx)

    assert len(result.sections) == 1
    assert result.sections[0].page_number is None
    assert result.sections[0].text == (
        "Introduction paragraph\nSecond paragraph\nName | Score\nAda | 99"
    )
    assert result.page_count is None


def test_docx_reads_title(sample_docx: bytes) -> None:
    assert extract_docx(sample_docx).title == "Doc Title"


def test_docx_skips_blank_paragraphs() -> None:
    result = extract_docx(make_docx(["first", "   ", "second"]))

    assert result.sections[0].text == "first\nsecond"


def test_docx_merged_cells_are_not_repeated() -> None:
    document = docx.Document()
    table = document.add_table(rows=1, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "merged"
    table.cell(0, 2).text = "single"
    buffer = io.BytesIO()
    document.save(buffer)

    assert extract_docx(buffer.getvalue()).sections[0].text == "merged | single"


def test_docx_damaged_raises_parse_error() -> None:
    with pytest.raises(DocumentParseError, match="damaged"):
        extract_docx(b"PK\x03\x04 not a real zip archive")


# --- TXT ---


def test_txt_decodes_utf8() -> None:
    result = extract_txt("Héllo wörld".encode())

    assert result.sections[0].text == "Héllo wörld"
    assert result.sections[0].page_number is None


def test_txt_strips_byte_order_mark() -> None:
    result = extract_txt(b"\xef\xbb\xbfwith BOM")

    assert result.sections[0].text == "with BOM"


def test_txt_normalizes_windows_line_endings() -> None:
    assert extract_txt(b"one\r\ntwo").sections[0].text == "one\ntwo"


def test_txt_rejects_non_utf8() -> None:
    with pytest.raises(DocumentParseError, match="UTF-8"):
        extract_txt("café".encode("latin-1"))


def test_txt_rejects_binary_content() -> None:
    with pytest.raises(DocumentParseError, match="binary"):
        extract_txt(b"abc\x00def")
