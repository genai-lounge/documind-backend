"""Integration tests: the DOCX adapter against the real python-docx library."""

import io

import docx
import pytest

from app.domain.errors import DocumentParseError
from app.infrastructure.parsing.python_docx_parser import PythonDocxParser
from tests.factories import make_docx

parser = PythonDocxParser()


def test_extracts_paragraphs_and_table_in_order(sample_docx: bytes) -> None:
    result = parser.parse(sample_docx)

    assert len(result.sections) == 1
    assert result.sections[0].page_number is None
    assert result.sections[0].text == (
        "Introduction paragraph\nSecond paragraph\nName | Score\nAda | 99"
    )
    assert result.page_count is None


def test_reads_title(sample_docx: bytes) -> None:
    assert parser.parse(sample_docx).metadata.title == "Doc Title"


def test_skips_blank_paragraphs() -> None:
    result = parser.parse(make_docx(["first", "   ", "second"]))

    assert result.sections[0].text == "first\nsecond"


def test_merged_cells_are_not_repeated() -> None:
    document = docx.Document()
    table = document.add_table(rows=1, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "merged"
    table.cell(0, 2).text = "single"
    buffer = io.BytesIO()
    document.save(buffer)

    assert parser.parse(buffer.getvalue()).sections[0].text == "merged | single"


def test_damaged_raises_parse_error() -> None:
    with pytest.raises(DocumentParseError, match="damaged"):
        parser.parse(b"PK\x03\x04 not a real zip archive")


@pytest.mark.parametrize(
    ("data", "expected"), [(b"PK\x03\x04rest", True), (b"just text", False)]
)
def test_matches_content(data: bytes, expected: bool) -> None:
    assert parser.matches_content(data) is expected
