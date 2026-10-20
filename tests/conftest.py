"""Shared test fixtures that build small, deterministic documents in memory.

Generating the files in code keeps binary test files out of the repository
and makes the content of each test document obvious.
"""

import io

import docx
import pymupdf
import pytest


def make_pdf(
    page_texts: list[str],
    title: str = "",
    author: str = "",
    password: str | None = None,
) -> bytes:
    """Build a PDF with one page per entry; an empty string makes a page with no text."""
    pdf = pymupdf.open()
    for text in page_texts:
        page = pdf.new_page()
        if text:
            page.insert_text((72, 72), text)
    pdf.set_metadata({"title": title, "author": author})
    if password:
        return pdf.tobytes(
            encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw=password, owner_pw=password
        )
    return pdf.tobytes()


def make_docx(
    paragraphs: list[str], table: list[list[str]] | None = None, title: str = ""
) -> bytes:
    """Build a DOCX with the given paragraphs, then an optional table."""
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        word_table = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                word_table.cell(r, c).text = value
    document.core_properties.title = title
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def two_page_pdf() -> bytes:
    return make_pdf(
        ["Alpha page one", "Beta page two"], title="Sample", author="Tester"
    )


@pytest.fixture
def sample_docx() -> bytes:
    return make_docx(
        ["Introduction paragraph", "Second paragraph"],
        table=[["Name", "Score"], ["Ada", "99"]],
        title="Doc Title",
    )
